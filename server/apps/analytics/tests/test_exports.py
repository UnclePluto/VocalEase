from io import BytesIO
import tempfile

from openpyxl import load_workbook

from apps.analytics.exporters import CsvExporter, XlsxExporter, export_rows_to


ROWS = [
    {
        "medical_record_no": "P000001",
        "name": "  =HYPERLINK(\"https://invalid\")",
        "treatment_status": "进行中",
        "primary_doctor": "王医生",
        "completed_count": 3,
        "total_duration_seconds": 180,
        "average_score": "88.50",
        "score_trend": "+2.00",
        "burp_improvement": "25.00%",
        "is_mock": "是",
    }
]


def test_csv_is_bom_encoded_and_escapes_formula_after_whitespace():
    content = CsvExporter().render(ROWS)
    decoded = content.decode("utf-8-sig")
    assert content.startswith(b"\xef\xbb\xbf")
    assert "'  =HYPERLINK" in decoded
    assert decoded.splitlines()[0] == "病历号,患者姓名,治疗状态,主治医生,完成演唱次数,累计时长（秒）,平均得分,得分趋势,嗳气改善率,模拟数据"


def test_xlsx_can_be_read_back_and_escapes_formula_cells():
    content = XlsxExporter().render(ROWS)
    sheet = load_workbook(BytesIO(content), read_only=True).active
    assert sheet.cell(1, 1).value == "病历号"
    assert sheet.cell(2, 2).value == "'  =HYPERLINK(\"https://invalid\")"
    assert sheet.cell(2, 2).data_type == "s"


def test_exporters_replace_illegal_controls_and_bound_long_cells():
    rows = [{**ROWS[0], "name": "安全\x00姓名" + "x" * 40000}]
    csv_value = CsvExporter().render(rows).decode("utf-8-sig")
    assert "\x00" not in csv_value
    sheet = load_workbook(BytesIO(XlsxExporter().render(rows)), read_only=True).active
    assert "\x00" not in sheet.cell(2, 2).value
    assert len(sheet.cell(2, 2).value) == 32767


def test_streaming_xlsx_large_batch_can_be_read_back():
    rows = ({**ROWS[0], "medical_record_no": f"P{index:06d}"} for index in range(1205))
    with tempfile.SpooledTemporaryFile(max_size=1024, mode="w+b") as stream:
        export_rows_to(stream, rows, "xlsx")
        stream.seek(0)
        sheet = load_workbook(stream, read_only=True).active
        values = list(sheet.iter_rows(values_only=True))

    assert len(values) == 1206
    assert values[1][0] == "P000000" and values[-1][0] == "P001204"


def test_xlsx_row_failure_closes_writer_and_removes_openpyxl_tempfile():
    from openpyxl.worksheet._writer import ALL_TEMP_FILES

    before = set(ALL_TEMP_FILES)

    def rows():
        yield ROWS[0]
        raise RuntimeError("row-failed")

    with tempfile.SpooledTemporaryFile(mode="w+b") as stream:
        try:
            export_rows_to(stream, rows(), "xlsx")
        except RuntimeError as exc:
            assert str(exc) == "row-failed"

    assert set(ALL_TEMP_FILES) == before


def test_xlsx_save_failure_closes_writer_and_removes_openpyxl_tempfile(monkeypatch):
    from openpyxl import Workbook
    from openpyxl.worksheet._writer import ALL_TEMP_FILES

    before = set(ALL_TEMP_FILES)
    monkeypatch.setattr(
        Workbook,
        "save",
        lambda self, stream: (_ for _ in ()).throw(RuntimeError("save-failed")),
    )
    with tempfile.SpooledTemporaryFile(mode="w+b") as stream:
        try:
            export_rows_to(stream, ROWS, "xlsx")
        except RuntimeError as exc:
            assert str(exc) == "save-failed"

    assert set(ALL_TEMP_FILES) == before
