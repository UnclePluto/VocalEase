from io import BytesIO

import pytest
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APIClient

from apps.analytics.models import ExportJob, ExportJobItem
from apps.analytics.tasks import expire_export_job, recover_export_jobs, run_export_job
from apps.audit.models import AuditLog
from apps.media.models import MediaAsset
from apps.accounts.models import Role, User
from apps.patients.models import PatientProfile, TreatmentPlan


@pytest.mark.django_db
def test_sync_export_intersects_selected_ids_with_current_filters(doctor, patient, other_patient):
    doctor.user.must_change_password = False
    doctor.user.save(update_fields=["must_change_password"])
    client = APIClient(); client.force_authenticate(doctor.user)
    response = client.post("/api/v1/admin/analytics/exports/", {
        "format": "xlsx",
        "filters": {"name": "患者甲"},
        "selected_ids": [str(patient.id), str(other_patient.id)],
        "idempotency_key": "sync-filter-intersection",
    }, format="json")

    assert response.status_code == 200
    assert response["Content-Type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "filename*=UTF-8''" in response["Content-Disposition"]
    rows = list(load_workbook(BytesIO(response.content), read_only=True).active.iter_rows(values_only=True))
    assert len(rows) == 2 and rows[1][0] == patient.medical_record_no
    assert AuditLog.objects.filter(action="analytics.export_sync", actor=doctor.user).count() == 1
    assert MediaAsset.objects.filter(owner_type="export").count() == 0


@pytest.mark.django_db
def test_async_export_snapshots_only_ids_and_defers_metrics_to_worker(
    settings, doctor, patient, other_patient, tmp_path, monkeypatch,
):
    settings.ANALYTICS_SYNC_EXPORT_LIMIT = 1
    settings.MEDIA_LOCAL_ROOT = str(tmp_path)
    doctor.user.must_change_password = False
    doctor.user.save(update_fields=["must_change_password"])
    client = APIClient(); client.force_authenticate(doctor.user)
    original_metric_rows = __import__("apps.analytics.views", fromlist=["patient_metric_rows"]).patient_metric_rows
    metric_calls = []

    def metric_rows_only_in_worker(patients):
        metric_calls.append(tuple(str(patient.id) for patient in patients))
        return original_metric_rows(patients)

    monkeypatch.setattr("apps.analytics.views.patient_metric_rows", lambda patients: pytest.fail("大导出请求不得计算指标"))
    response = client.post("/api/v1/admin/analytics/exports/", {
        "format": "csv",
        "filters": {"treatment_status": "active"},
        "selected_ids": [],
        "idempotency_key": "async-active-patients",
    }, format="json")

    assert response.status_code == 202
    job = ExportJob.objects.get(pk=response.json()["data"]["id"])
    assert job.snapshot_count == 2 and job.normalized_filters["treatment_status"] == "active"
    assert not hasattr(job, "rows_snapshot")
    assert list(ExportJobItem.objects.filter(job=job).order_by("position").values_list("patient_id", flat=True)) == [
        patient.id, other_patient.id,
    ]

    # 创建后的软删不能改变该 Job 已固定的患者集合。
    type(patient).objects.filter(pk=patient.id).update(deleted_at=timezone.now())
    monkeypatch.setattr("apps.analytics.tasks.patient_metric_rows", metric_rows_only_in_worker)
    run_export_job(str(job.id))

    job.refresh_from_db()
    asset = MediaAsset.objects.get(pk=job.result_asset_id)
    assert job.status == "ready" and asset.status == "ready"
    assert asset.owner_type == "export" and asset.owner_id == job.id
    assert {str(patient.id), str(other_patient.id)} == set().union(*map(set, metric_calls))
    private = client.post(f"/api/v1/admin/analytics/exports/{job.id}/private-url/")
    assert private.status_code == 200
    assert private.json()["data"]["expires_at"] <= job.expires_at.isoformat()
    assert AuditLog.objects.filter(action="analytics.export_create", target_id=job.id).exists()
    assert AuditLog.objects.filter(action="analytics.export_download", target_id=job.id).exists()


@pytest.mark.django_db
def test_export_idempotency_rejects_changed_request(settings, doctor, patient, other_patient):
    settings.ANALYTICS_SYNC_EXPORT_LIMIT = 1
    doctor.user.must_change_password = False
    doctor.user.save(update_fields=["must_change_password"])
    client = APIClient(); client.force_authenticate(doctor.user)
    payload = {"format": "csv", "filters": {}, "selected_ids": [], "idempotency_key": "same-key"}
    first = client.post("/api/v1/admin/analytics/exports/", payload, format="json")
    second = client.post("/api/v1/admin/analytics/exports/", payload, format="json")
    changed = client.post("/api/v1/admin/analytics/exports/", {**payload, "format": "xlsx"}, format="json")

    assert first.status_code == 202 and second.status_code == 200
    assert first.json()["data"]["id"] == second.json()["data"]["id"]
    assert changed.status_code == 409
    assert ExportJob.objects.count() == 1


@pytest.mark.django_db
def test_more_than_1000_rows_never_builds_request_metrics_or_json_snapshot(doctor, patient, monkeypatch):
    users = [
        User(login_id=f"bulk-patient-{index:04d}", password="!", role=Role.PATIENT, must_change_password=False)
        for index in range(1000)
    ]
    User.objects.bulk_create(users, batch_size=250)
    patients = [
        PatientProfile(
            user=user,
            medical_record_no=f"PX{index:06d}",
            name=f"批量患者{index:04d}",
            gender="male",
            enrollment_age=30,
            phone=f"139{index:08d}",
            primary_doctor=doctor,
        )
        for index, user in enumerate(users)
    ]
    PatientProfile.objects.bulk_create(patients, batch_size=250)
    TreatmentPlan.objects.bulk_create([
        TreatmentPlan(
            patient=item, start_date="2026-08-01", cycle_weeks=4,
            target_session_count=12, status=TreatmentPlan.Status.ACTIVE,
        )
        for item in patients
    ], batch_size=250)
    monkeypatch.setattr(
        "apps.analytics.views.patient_metric_rows",
        lambda patients: pytest.fail(">1000 请求端不得构造指标行"),
    )
    client = APIClient(); client.force_authenticate(doctor.user)

    response = client.post("/api/v1/admin/analytics/exports/", {
        "format": "csv", "filters": {}, "selected_ids": [], "idempotency_key": "real-over-1000",
    }, format="json")

    assert response.status_code == 202
    job = ExportJob.objects.get(pk=response.json()["data"]["id"])
    assert job.snapshot_count == 1001
    assert ExportJobItem.objects.filter(job=job).count() == 1001
    assert "rows_snapshot" not in {field.name for field in ExportJob._meta.fields}


@pytest.mark.django_db
def test_patient_cannot_read_or_download_export_job(settings, doctor, patient, other_patient):
    settings.ANALYTICS_SYNC_EXPORT_LIMIT = 1
    doctor.user.must_change_password = False
    doctor.user.save(update_fields=["must_change_password"])
    client = APIClient(); client.force_authenticate(doctor.user)
    created = client.post("/api/v1/admin/analytics/exports/", {
        "format": "csv", "filters": {}, "selected_ids": [], "idempotency_key": "patient-denied",
    }, format="json")
    job_id = created.json()["data"]["id"]
    client.force_authenticate(patient.user)
    assert client.get(f"/api/v1/admin/analytics/exports/{job_id}/").status_code == 403
    assert client.post(f"/api/v1/admin/analytics/exports/{job_id}/private-url/").status_code == 403


@pytest.mark.django_db
def test_export_failure_never_leaves_false_ready_job(settings, doctor, patient, other_patient, monkeypatch):
    settings.ANALYTICS_SYNC_EXPORT_LIMIT = 1
    doctor.user.must_change_password = False
    doctor.user.save(update_fields=["must_change_password"])
    client = APIClient(); client.force_authenticate(doctor.user)
    created = client.post("/api/v1/admin/analytics/exports/", {
        "format": "csv", "filters": {}, "selected_ids": [], "idempotency_key": "publish-crash",
    }, format="json")
    job = ExportJob.objects.get(pk=created.json()["data"]["id"])
    monkeypatch.setattr("apps.analytics.tasks.publish_generated_asset", lambda **kwargs: (_ for _ in ()).throw(RuntimeError("storage-down")))

    with pytest.raises(RuntimeError, match="storage-down"):
        run_export_job(str(job.id))

    job.refresh_from_db()
    assert job.status == "pending" and job.result_asset_id is None
    assert job.claim_token is None and job.lease_expires_at is None


@pytest.mark.django_db
def test_recovery_requeues_expired_lease_without_double_claim(settings, doctor, monkeypatch):
    from datetime import timedelta
    import uuid

    job = ExportJob.objects.create(
        creator=doctor.user, normalized_filters={}, selected_ids=[], snapshot_count=0,
        format="csv", status="processing", expires_at=timezone.now() + timedelta(hours=1),
        idempotency_key="recover-expired", request_fingerprint="a" * 64, attempt=1,
        claim_token=uuid.uuid4(), heartbeat_at=timezone.now() - timedelta(minutes=10),
        lease_expires_at=timezone.now() - timedelta(minutes=5),
    )
    dispatched = []
    monkeypatch.setattr("apps.analytics.tasks.run_export_job_task.delay", lambda job_id: dispatched.append(job_id))

    result = recover_export_jobs()

    job.refresh_from_db()
    assert result == {"recovered": 1, "failed": 0, "expired": 0, "cleaned": 0}
    assert job.status == "pending" and job.claim_token is None
    assert dispatched == [str(job.id)]


@pytest.mark.django_db
def test_dashboard_rejects_unknown_query_parameters(doctor):
    client = APIClient(); client.force_authenticate(doctor.user)
    response = client.get("/api/v1/admin/analytics/dashboard/?unexpected=1")
    assert response.status_code == 400
    assert "unexpected" in response.json()["data"]


@pytest.mark.django_db
def test_export_worker_renews_claim_after_render_before_publishing(settings, doctor, monkeypatch):
    from datetime import timedelta
    import uuid

    job = ExportJob.objects.create(
        creator=doctor.user, normalized_filters={}, selected_ids=[], snapshot_count=0,
        format="csv", expires_at=timezone.now() + timedelta(hours=1), idempotency_key="heartbeat-before-publish",
        request_fingerprint="c" * 64,
    )
    asset = MediaAsset.objects.create(
        owner_type="export", owner_id=job.id, media_type="export", backend="qiniu",
        object_key=f"test/export/{uuid.uuid4().hex}", mime="text/csv", size=3,
        etag="export-ready", status="ready", upload_expires_at=timezone.now() + timedelta(hours=1),
    )

    def slow_render(stream, rows, export_format):
        stream.write(b"abc")
        ExportJob.objects.filter(pk=job.id).update(lease_expires_at=timezone.now() + timedelta(seconds=1))

    def publish(**kwargs):
        leased = ExportJob.objects.get(pk=job.id)
        assert leased.lease_expires_at > timezone.now() + timedelta(seconds=100)
        return asset

    monkeypatch.setattr("apps.analytics.tasks.export_rows_to", slow_render)
    monkeypatch.setattr("apps.analytics.tasks.publish_generated_asset", publish)
    monkeypatch.setattr("apps.analytics.tasks.resolve_export_asset", lambda job, verify_storage=True: asset)

    result = run_export_job(str(job.id))

    job.refresh_from_db()
    assert result["published"] is True
    assert job.status == "ready" and job.result_asset_id == asset.id


@pytest.mark.django_db
def test_expired_export_revokes_its_private_media(doctor):
    from datetime import timedelta
    import uuid

    job = ExportJob.objects.create(
        creator=doctor.user, normalized_filters={}, selected_ids=[], snapshot_count=0,
        format="csv", status="pending", expires_at=timezone.now() - timedelta(seconds=1),
        idempotency_key="expire-media", request_fingerprint="d" * 64,
    )
    asset = MediaAsset.objects.create(
        owner_type="export", owner_id=job.id, media_type="export", backend="qiniu",
        object_key=f"test/export/{uuid.uuid4().hex}", mime="text/csv", size=3,
        etag="expired-ready", status="ready", upload_expires_at=timezone.now() + timedelta(hours=1),
    )
    ExportJob.objects.filter(pk=job.id).update(status="ready", result_asset_id=asset.id, completed_at=timezone.now())
    client = APIClient(); client.force_authenticate(doctor.user)

    response = client.get(f"/api/v1/admin/analytics/exports/{job.id}/")

    job.refresh_from_db(); asset.refresh_from_db()
    assert response.status_code == 200 and job.status == "expired"
    assert asset.status == "pending_cleanup"


@pytest.mark.django_db
def test_worker_never_links_media_revoked_during_lease_takeover(doctor, monkeypatch):
    from datetime import timedelta
    import uuid

    job = ExportJob.objects.create(
        creator=doctor.user, normalized_filters={}, selected_ids=[], snapshot_count=0,
        format="csv", expires_at=timezone.now() + timedelta(hours=1),
        idempotency_key="revoked-before-finalize", request_fingerprint="e" * 64,
    )
    asset = MediaAsset.objects.create(
        owner_type="export", owner_id=job.id, media_type="export", backend="qiniu",
        object_key=f"test/export/{uuid.uuid4().hex}", mime="text/csv", size=3,
        etag="revoked", status="pending_cleanup", upload_expires_at=timezone.now() + timedelta(hours=1),
    )
    monkeypatch.setattr("apps.analytics.tasks.export_rows_to", lambda stream, rows, export_format: stream.write(b"abc"))
    monkeypatch.setattr("apps.analytics.tasks.publish_generated_asset", lambda **kwargs: asset)

    result = run_export_job(str(job.id))

    job.refresh_from_db()
    assert result["published"] is False
    assert job.status == "pending" and job.result_asset_id is None


@pytest.mark.django_db
def test_worker_does_not_publish_ready_after_job_expires(doctor, monkeypatch):
    from datetime import timedelta
    import uuid

    job = ExportJob.objects.create(
        creator=doctor.user, normalized_filters={}, selected_ids=[], snapshot_count=0,
        format="csv", expires_at=timezone.now() + timedelta(minutes=1),
        idempotency_key="expire-during-publish", request_fingerprint="f" * 64,
    )
    asset = MediaAsset.objects.create(
        owner_type="export", owner_id=job.id, media_type="export", backend="qiniu",
        object_key=f"test/export/{uuid.uuid4().hex}", mime="text/csv", size=3,
        etag="late-ready", status="ready", upload_expires_at=timezone.now() + timedelta(hours=1),
    )

    def publish(**kwargs):
        ExportJob.objects.filter(pk=job.id).update(expires_at=timezone.now() - timedelta(seconds=1))
        return asset

    monkeypatch.setattr("apps.analytics.tasks.export_rows_to", lambda stream, rows, export_format: stream.write(b"abc"))
    monkeypatch.setattr("apps.analytics.tasks.publish_generated_asset", publish)

    result = run_export_job(str(job.id))

    job.refresh_from_db(); asset.refresh_from_db()
    assert result["published"] is False and job.status == "expired"
    assert asset.status == "pending_cleanup"


@pytest.mark.django_db
def test_cleanup_crash_is_recovered_and_clears_snapshot_once(settings, doctor, patient, monkeypatch):
    from datetime import timedelta
    import uuid

    job = ExportJob.objects.create(
        creator=doctor.user,
        normalized_filters={"name": "敏感患者"},
        selected_ids=[str(patient.id)],
        snapshot_count=1,
        format="csv",
        expires_at=timezone.now() - timedelta(seconds=1),
        idempotency_key="cleanup-crash",
        request_fingerprint="1" * 64,
    )
    ExportJobItem.objects.create(job=job, patient_id=patient.id, position=0)
    asset = MediaAsset.objects.create(
        owner_type="export", owner_id=job.id, media_type="export", backend="qiniu",
        object_key=f"test/export/{uuid.uuid4().hex}", mime="text/csv", size=3,
        etag="cleanup-crash", status="ready", upload_expires_at=timezone.now() + timedelta(hours=1),
    )
    ExportJob.objects.filter(pk=job.id).update(
        status="ready", result_asset_id=asset.id, completed_at=timezone.now(),
    )
    original = __import__("apps.analytics.tasks", fromlist=["mark_asset_for_cleanup"]).mark_asset_for_cleanup
    monkeypatch.setattr(
        "apps.analytics.tasks.mark_asset_for_cleanup",
        lambda **kwargs: (_ for _ in ()).throw(RuntimeError("cleanup-crash")),
    )

    # GET/定时过期即便在媒体标记点崩溃也不能 500，任务保留可恢复清理租约。
    expire_export_job(job.id)
    job.refresh_from_db()
    assert job.status == "expired" and job.cleanup_status == "processing"
    assert not ExportJobItem.objects.filter(job=job).exists()

    monkeypatch.setattr("apps.analytics.tasks.mark_asset_for_cleanup", original)
    ExportJob.objects.filter(pk=job.id).update(cleanup_lease_expires_at=timezone.now() - timedelta(seconds=1))
    recovered = recover_export_jobs()

    job.refresh_from_db(); asset.refresh_from_db()
    assert recovered["cleaned"] == 1
    assert job.cleanup_status == "complete" and job.result_asset_id is None
    assert job.normalized_filters == {} and job.selected_ids == []
    assert asset.status == "pending_cleanup"
    assert AuditLog.objects.filter(action="analytics.export_cleanup", target_id=job.id).count() == 1


@pytest.mark.django_db
@pytest.mark.parametrize("asset_variant", ["missing", "soft_deleted", "wrong_mime", "wrong_owner"])
def test_private_url_rejects_untrusted_or_missing_result_without_500(doctor, asset_variant):
    from datetime import timedelta
    import uuid

    job = ExportJob.objects.create(
        creator=doctor.user, normalized_filters={}, selected_ids=[], snapshot_count=0,
        format="csv", expires_at=timezone.now() + timedelta(hours=1),
        idempotency_key=f"invalid-result-{asset_variant}", request_fingerprint="2" * 64,
    )
    asset_id = uuid.uuid4()
    if asset_variant != "missing":
        asset = MediaAsset.objects.create(
            owner_type="export",
            owner_id=uuid.uuid4() if asset_variant == "wrong_owner" else job.id,
            media_type="export", backend="qiniu",
            object_key=f"test/export/{uuid.uuid4().hex}",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" if asset_variant == "wrong_mime" else "text/csv",
            size=3, etag="invalid-result", status="ready",
            upload_expires_at=timezone.now() + timedelta(hours=1),
        )
        asset_id = asset.id
        if asset_variant == "soft_deleted":
            MediaAsset.objects.filter(pk=asset.id).update(deleted_at=timezone.now())
    ExportJob.objects.filter(pk=job.id).update(
        status="ready", result_asset_id=asset_id, completed_at=timezone.now(),
    )
    doctor.user.must_change_password = False
    doctor.user.save(update_fields=["must_change_password"])
    client = APIClient(); client.force_authenticate(doctor.user)

    response = client.post(f"/api/v1/admin/analytics/exports/{job.id}/private-url/")

    job.refresh_from_db()
    assert response.status_code == 409
    assert response.json()["code"] in {
        "export_asset_missing", "export_asset_invalid",
    }
    assert job.status == "failed" and job.cleanup_status == "complete"


@pytest.mark.django_db
def test_private_url_rejects_local_manifest_receipt_mismatch(
    settings, doctor, patient, other_patient, tmp_path,
):
    settings.ANALYTICS_SYNC_EXPORT_LIMIT = 1
    settings.MEDIA_BACKEND = "local"
    settings.MEDIA_LOCAL_ROOT = str(tmp_path)
    doctor.user.must_change_password = False
    doctor.user.save(update_fields=["must_change_password"])
    client = APIClient(); client.force_authenticate(doctor.user)
    created = client.post("/api/v1/admin/analytics/exports/", {
        "format": "csv", "filters": {}, "selected_ids": [], "idempotency_key": "receipt-mismatch",
    }, format="json")
    job = ExportJob.objects.get(pk=created.json()["data"]["id"])
    run_export_job(str(job.id))
    job.refresh_from_db()
    MediaAsset.objects.filter(pk=job.result_asset_id).update(sha256="0" * 64)

    response = client.post(f"/api/v1/admin/analytics/exports/{job.id}/private-url/")

    job.refresh_from_db()
    assert response.status_code == 409
    assert response.json()["code"] in {"export_asset_invalid", "export_asset_unverifiable"}
    assert job.status == "failed" and job.cleanup_status == "complete"
