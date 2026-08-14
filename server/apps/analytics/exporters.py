from __future__ import annotations

import csv
from io import BytesIO, StringIO, TextIOWrapper
import os
import re
from typing import Iterable, Mapping

from openpyxl import Workbook


EXPORT_COLUMNS = (
    ("medical_record_no", "病历号"),
    ("name", "患者姓名"),
    ("treatment_status", "治疗状态"),
    ("primary_doctor", "主治医生"),
    ("completed_count", "完成演唱次数"),
    ("total_duration_seconds", "累计时长（秒）"),
    ("average_score", "平均得分"),
    ("score_trend", "得分趋势"),
    ("burp_improvement", "嗳气改善率"),
    ("is_mock", "模拟数据"),
)
MAX_CELL_LENGTH = 32767
_ILLEGAL_CONTROLS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")


def escape_sheet_cell(value):
    if value is None:
        return ""
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return value
    text = _ILLEGAL_CONTROLS.sub("�", str(value))[:MAX_CELL_LENGTH]
    effective = text.lstrip()
    if effective.startswith(("=", "+", "-", "@")):
        text = "'" + text
        text = text[:MAX_CELL_LENGTH]
    return text


class CsvExporter:
    mime = "text/csv"
    extension = "csv"

    def render(self, rows: Iterable[Mapping[str, object]]) -> bytes:
        stream = StringIO(newline="")
        writer = csv.writer(stream, lineterminator="\r\n")
        writer.writerow([label for _, label in EXPORT_COLUMNS])
        for row in rows:
            writer.writerow([escape_sheet_cell(row.get(key)) for key, _ in EXPORT_COLUMNS])
        return b"\xef\xbb\xbf" + stream.getvalue().encode("utf-8")

    def render_to(self, stream, rows: Iterable[Mapping[str, object]], *, heartbeat=None) -> None:
        stream.write(b"\xef\xbb\xbf")
        text_stream = TextIOWrapper(stream, encoding="utf-8", newline="", write_through=True)
        try:
            writer = csv.writer(text_stream, lineterminator="\r\n")
            writer.writerow([label for _, label in EXPORT_COLUMNS])
            for row in rows:
                if heartbeat:
                    heartbeat()
                writer.writerow([escape_sheet_cell(row.get(key)) for key, _ in EXPORT_COLUMNS])
            text_stream.flush()
            if heartbeat:
                heartbeat()
        finally:
            text_stream.detach()


class XlsxExporter:
    mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    extension = "xlsx"

    def render(self, rows: Iterable[Mapping[str, object]]) -> bytes:
        workbook = Workbook(write_only=True)
        stream = BytesIO()
        try:
            worksheet = workbook.create_sheet("患者统计")
            worksheet.append([label for _, label in EXPORT_COLUMNS])
            for row in rows:
                worksheet.append([escape_sheet_cell(row.get(key)) for key, _ in EXPORT_COLUMNS])
            workbook.save(stream)
            return stream.getvalue()
        finally:
            _close_write_only_workbook(workbook)

    def render_to(self, stream, rows: Iterable[Mapping[str, object]], *, heartbeat=None) -> None:
        workbook = Workbook(write_only=True)
        try:
            worksheet = workbook.create_sheet("患者统计")
            worksheet.append([label for _, label in EXPORT_COLUMNS])
            for row in rows:
                if heartbeat:
                    heartbeat()
                worksheet.append([escape_sheet_cell(row.get(key)) for key, _ in EXPORT_COLUMNS])
            workbook.save(_HeartbeatWriteProxy(stream, heartbeat) if heartbeat else stream)
        finally:
            _close_write_only_workbook(workbook)


def _close_write_only_workbook(workbook) -> None:
    """显式关闭 openpyxl 3.1 write-only writer 并移除其命名临时文件。"""
    for worksheet in workbook.worksheets:
        rows = getattr(worksheet, "_rows", None)
        writer = getattr(worksheet, "_writer", None)
        if rows is not None:
            try:
                rows.close()
            except Exception:
                pass
        if writer is not None:
            try:
                writer.close()
            except Exception:
                pass
            if os.path.exists(writer.out):
                try:
                    writer.cleanup()
                except (OSError, ValueError):
                    pass
    workbook.close()


class _HeartbeatWriteProxy:
    def __init__(self, stream, heartbeat):
        self._stream = stream
        self._heartbeat = heartbeat

    def write(self, data):
        self._heartbeat()
        written = self._stream.write(data)
        self._heartbeat()
        return written

    def __getattr__(self, name):
        return getattr(self._stream, name)


def export_rows(rows: Iterable[Mapping[str, object]], export_format: str) -> bytes:
    exporter = CsvExporter() if export_format == "csv" else XlsxExporter() if export_format == "xlsx" else None
    if exporter is None:
        raise ValueError("不支持的导出格式")
    return exporter.render(rows)


def export_rows_to(stream, rows: Iterable[Mapping[str, object]], export_format: str, *, heartbeat=None) -> None:
    exporter = CsvExporter() if export_format == "csv" else XlsxExporter() if export_format == "xlsx" else None
    if exporter is None:
        raise ValueError("不支持的导出格式")
    exporter.render_to(stream, rows, heartbeat=heartbeat)
